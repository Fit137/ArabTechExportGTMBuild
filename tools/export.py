#!/usr/bin/env python3
"""Generate the XLSX and DOCX deliverables the gtm-builder skill's
OUTPUT FORMAT MAPPING calls for, from the Markdown source of truth.

Markdown under deliverables/ is authoritative. Everything in
deliverables/exports/ is generated and safe to delete. Re-run after any edit:

    python3 tools/export.py

Usage: no arguments. Paths are resolved relative to the repo root.
"""

from __future__ import annotations

import re
import shutil
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches
from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "deliverables"
OUT = SRC / "exports"

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(size=10)
TBD_FONT = Font(size=10, bold=True, color="B45309")
THIN = Side(style="thin", color="D0D7E5")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

NUMERIC = re.compile(r"^-?\d+(?:\.\d+)?$")


def coerce(value: str):
    """Store bare numbers as numbers, so charts and sorting work in Excel."""
    return float(value) if NUMERIC.fullmatch(value) else value


# ---------------------------------------------------------------- parsing


@dataclass
class Block:
    """One parsed chunk of Markdown: a heading, a paragraph, or a table."""

    kind: str  # heading | para | table | bullet | code
    level: int = 0
    text: str = ""
    rows: list[list[str]] = field(default_factory=list)


def strip_md(cell: str) -> str:
    """Reduce inline Markdown to plain text, preserving the visible words."""
    cell = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", cell)  # links
    cell = re.sub(r"(\*\*|\*|`|~~)", "", cell)
    return cell.replace(r"\|", "|").strip()


def is_table_row(line: str) -> bool:
    return line.lstrip().startswith("|") and line.rstrip().endswith("|")


def is_separator_row(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def split_row(line: str) -> list[str]:
    return [strip_md(c) for c in line.strip().strip("|").split("|")]


def parse(md: str) -> list[Block]:
    blocks: list[Block] = []
    lines = md.splitlines()
    i, in_fence = 0, False
    para: list[str] = []

    def flush_para() -> None:
        if para:
            blocks.append(Block("para", text=" ".join(para).strip()))
            para.clear()

    while i < len(lines):
        line = lines[i]

        if line.lstrip().startswith("```"):
            flush_para()
            in_fence = not in_fence
            if in_fence:
                buf, i = [], i + 1
                while i < len(lines) and not lines[i].lstrip().startswith("```"):
                    buf.append(lines[i])
                    i += 1
                in_fence = False
                blocks.append(Block("code", text="\n".join(buf)))
            i += 1
            continue

        if not line.strip():
            flush_para()
            i += 1
            continue

        if m := re.match(r"^(#{1,6})\s+(.*)", line):
            flush_para()
            blocks.append(Block("heading", level=len(m.group(1)), text=strip_md(m.group(2))))
            i += 1
            continue

        if m := re.match(r"^\s*(?:[-*]|\d+\.)\s+(.*)", line):
            flush_para()
            blocks.append(Block("bullet", text=strip_md(m.group(1))))
            i += 1
            continue

        if is_table_row(line) and i + 1 < len(lines) and is_separator_row(lines[i + 1]):
            flush_para()
            rows = [split_row(line)]
            i += 2
            while i < len(lines) and is_table_row(lines[i]):
                rows.append(split_row(lines[i]))
                i += 1
            blocks.append(Block("table", rows=rows))
            continue

        para.append(line.strip())
        i += 1

    flush_para()
    return blocks


# ---------------------------------------------------------------- xlsx


def safe_sheet_name(name: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", name)
    name = re.sub(r"\s+", " ", name).strip() or "Sheet"
    if len(name) > 31:
        clipped = name[:31]
        if " " in clipped[20:]:
            clipped = clipped[: clipped.rindex(" ")]
        name = clipped.rstrip(" .,")
    base, n = name, 2
    while name.lower() in used:
        suffix = f" {n}"
        name = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(name.lower())
    return name


def sheet_label(heading: str) -> str:
    """Turn 'Step 2.10. Feature presence' into something readable in a tab."""
    heading = re.sub(r"^Steps?\s+", "", heading)
    heading = re.sub(r"^(\d+(?:\.\d+)*)[.)]?\s*", r"\1 ", heading)
    return heading.strip()


def write_table(ws, rows: list[list[str]], title: str | None, start: int) -> int:
    r = start
    if title:
        ws.cell(row=r, column=1, value=title).font = Font(bold=True, size=11, color="1F3864")
        r += 2

    header, *body = rows
    widths: dict[int, int] = {}

    for c, val in enumerate(header, 1):
        cell = ws.cell(row=r, column=c, value=val)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        widths[c] = len(val)
    r += 1

    for line in body:
        for c, val in enumerate(line, 1):
            cell = ws.cell(row=r, column=c, value=coerce(val))
            cell.font = TBD_FONT if "{TBD}" in val else BODY_FONT
            cell.border = BORDER
            cell.alignment = Alignment(
                vertical="top", wrap_text=True,
                horizontal="right" if isinstance(cell.value, float) else "left",
            )
            widths[c] = max(widths.get(c, 0), min(len(val), 70))
        r += 1

    for c, w in widths.items():
        letter = get_column_letter(c)
        current = ws.column_dimensions[letter].width or 0
        ws.column_dimensions[letter].width = max(current, min(max(w + 3, 12), 62))

    return r + 2


def add_scatter(ws, rows: list[list[str]], anchor_row: int) -> None:
    """Plot demand against competitiveness when a scoring table supplies both."""
    header = [h.lower() for h in rows[0]]
    try:
        xi = next(i for i, h in enumerate(header) if "demand" in h and "score" in h)
        yi = next(i for i, h in enumerate(header) if "competitive" in h and "score" in h)
    except StopIteration:
        return

    first = anchor_row + 3  # title, blank, header
    last = first + len(rows) - 2

    chart = ScatterChart()
    chart.title = "Demand against competitiveness"
    chart.style = 13
    chart.x_axis.title = "Demand"
    chart.y_axis.title = "Competitiveness"
    chart.height, chart.width = 11, 19

    xs = Reference(ws, min_col=xi + 1, min_row=first, max_row=last)
    ys = Reference(ws, min_col=yi + 1, min_row=first, max_row=last)
    series = Series(ys, xs, title="Service items")
    series.marker.symbol = "circle"
    series.marker.size = 9
    series.graphicalProperties.line.noFill = True
    chart.series.append(series)

    ws.add_chart(chart, f"A{last + 3}")


def to_xlsx(md_path: Path, out_path: Path, note: str) -> Path:
    blocks = parse(md_path.read_text())
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()

    # Sheet one: the prose, so context is not lost in the export.
    ws = wb.create_sheet(safe_sheet_name("About", used))
    ws.column_dimensions["A"].width = 110
    r = 1
    ws.cell(row=r, column=1, value=md_path.stem).font = Font(bold=True, size=14, color="1F3864")
    r += 1
    ws.cell(row=r, column=1, value=note).font = Font(size=9, italic=True, color="6B7280")
    r += 2
    for b in blocks:
        if b.kind == "heading":
            ws.cell(row=r, column=1, value=b.text).font = Font(bold=True, size=11, color="1F3864")
            r += 1
        elif b.kind in ("para", "code"):
            cell = ws.cell(row=r, column=1, value=b.text)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1
        elif b.kind == "bullet":
            cell = ws.cell(row=r, column=1, value=f"\u2022 {b.text}")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            r += 1

    # One sheet per table, named from the nearest preceding heading.
    heading, counts = "Table", {}
    for b in blocks:
        if b.kind == "heading" and b.level >= 2:
            heading = sheet_label(b.text)
        elif b.kind == "table":
            counts[heading] = counts.get(heading, 0) + 1
            label = heading if counts[heading] == 1 else f"{heading} ({counts[heading]})"
            ws = wb.create_sheet(safe_sheet_name(label, used))
            ws.freeze_panes = "A4"
            end = write_table(ws, b.rows, heading, 1)
            add_scatter(ws, b.rows, 1)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


# ---------------------------------------------------------------- docx


def to_docx(md_path: Path, out_path: Path, note: str, only_from: str | None = None) -> Path:
    blocks = parse(md_path.read_text())

    if only_from:
        for i, b in enumerate(blocks):
            if b.kind == "heading" and only_from.lower() in b.text.lower():
                blocks = blocks[i:]
                break

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    head = doc.add_paragraph()
    run = head.add_run(md_path.stem)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)

    sub = doc.add_paragraph()
    srun = sub.add_run(note)
    srun.italic = True
    srun.font.size = Pt(8.5)
    srun.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    for b in blocks:
        if b.kind == "heading":
            doc.add_heading(b.text, level=min(max(b.level, 1), 4))
        elif b.kind == "para":
            doc.add_paragraph(b.text)
        elif b.kind == "bullet":
            doc.add_paragraph(b.text, style="List Bullet")
        elif b.kind == "code":
            p = doc.add_paragraph()
            run = p.add_run(b.text)
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        elif b.kind == "table":
            cols = max(len(r) for r in b.rows)
            table = doc.add_table(rows=0, cols=cols)
            table.style = "Light Grid Accent 1"
            table.autofit = True
            for ri, row in enumerate(b.rows):
                cells = table.add_row().cells
                for ci in range(cols):
                    text = row[ci] if ci < len(row) else ""
                    para = cells[ci].paragraphs[0]
                    run = para.add_run(text)
                    run.font.size = Pt(8.5)
                    if ri == 0:
                        run.bold = True
                    elif "{TBD}" in text:
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)
            doc.add_paragraph()

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path


# ---------------------------------------------------------------- mapping

NOTE = (
    "Generated from the Markdown source in deliverables/ by tools/export.py. "
    "The Markdown is the source of truth. Edit there and re-run, do not edit this file."
)

# (source stem, formats). Formats: "xlsx", "docx", "docx-from:<heading fragment>"
MAPPING = [
    ("01-service-offer-matrix", ["xlsx"]),
    ("02-competitive-analysis", ["xlsx", "docx-from:Step 2.15"]),
    ("03-positioning", ["docx"]),
    ("04-icp", ["xlsx"]),
    ("05-buying-committee", ["xlsx"]),
    ("06-value-proposition", ["docx"]),
    ("07-customer-journey", ["xlsx"]),
    ("08-assets-library", ["xlsx"]),
    ("09-pricing-packaging", ["xlsx"]),
    ("10-lead-magnet-scorecard", ["xlsx", "docx"]),
    ("10-lead-magnet-diagnostic", ["xlsx", "docx"]),
]


def main() -> None:
    if OUT.exists():
        shutil.rmtree(OUT)

    made: list[Path] = []

    for track in ("track-b", "track-a"):
        for stem, formats in MAPPING:
            src = SRC / track / f"{stem}.md"
            if not src.exists():
                continue
            for fmt in formats:
                if fmt == "xlsx":
                    made.append(to_xlsx(src, OUT / track / f"{stem}.xlsx", NOTE))
                elif fmt == "docx":
                    made.append(to_docx(src, OUT / track / f"{stem}.docx", NOTE))
                elif fmt.startswith("docx-from:"):
                    heading = fmt.split(":", 1)[1]
                    made.append(
                        to_docx(
                            src,
                            OUT / track / f"{stem}-summary.docx",
                            NOTE + f" Contains the sections from '{heading}' onward.",
                            only_from=heading,
                        )
                    )

        # The deck stays Markdown. The skill specifies MD for Gamma AI.
        deck = SRC / track / "11-gtm-walkthrough-deck.md"
        if deck.exists():
            dest = OUT / track / "11-gtm-walkthrough-deck.md"
            dest.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(deck, dest)
            made.append(dest)

    # Not in the skill's mapping, but the competitor slate belongs in a sheet.
    made.append(to_xlsx(SRC / "00-RESEARCH-FINDINGS.md", OUT / "00-research-findings.xlsx", NOTE))
    made.append(to_xlsx(SRC / "00-BUYER-CHECKS.md", OUT / "00-buyer-checks.xlsx", NOTE))
    made.append(to_docx(SRC / "00-BUYER-CHECKS.md", OUT / "00-buyer-checks.docx", NOTE))
    made.append(to_docx(SRC / "00-CONFLICTS-TO-RESOLVE.md", OUT / "00-conflicts-to-resolve.docx", NOTE))
    made.append(to_docx(SRC / "00-RESOLUTIONS.md", OUT / "00-resolutions.docx", NOTE))

    for p in sorted(made):
        print(p.relative_to(ROOT))
    print(f"\n{len(made)} files written to {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
